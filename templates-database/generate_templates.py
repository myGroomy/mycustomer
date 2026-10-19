#!/usr/bin/env python3
"""
Generator template database (.xlsx) untuk MyCustomer.

Menghasilkan `mycustomer-database-template.xlsx` yang berisi semua tab yang
dibutuhkan aplikasi (lihat plan-database.md §3). Header dan urutan kolom harus
PERSIS dengan yang dibaca kode — beberapa bagian membaca kolom berdasarkan
index, bukan nama (sync-staging: customers A=id, B=phone, C=name;
orders B=customer_id, C=order_date, D=channel).

Cara pakai:
    python3 templates-database/generate_templates.py

Re-runnable: file .xlsx ditimpa setiap kali script dijalankan.
Skema berubah? Ubah SHEETS di bawah, jalankan ulang, selesai.
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = Path(__file__).parent / "mycustomer-database-template.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEED_FILL = PatternFill("solid", fgColor="E2EFDA")  # baris seed/default
SAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")  # baris contoh

# Kolom yang formatnya HARUS teks (nomor telepon/ID; cegah leading-zero hilang)
TEXT_COLUMNS = {
    "customers": ["id", "phone_normalized", "usia"],
    "orders": ["id", "customer_id", "raw_phone_input"],
    "branches": ["id"],
    "users": ["id", "pin"],
    "customer_branches": ["id", "customer_id", "local_name"],
    "staging_cmh": ["no_telp"],
    "staging_bdg": ["no_telp"],
}

CHANNEL_ENUM = "dine_in,takeaway,gofood,grab,shopee,whatsapp,custom"
BRANCH_ENUM = "CMH,BDG"
ROLE_ENUM = "owner,admin,kasir"
GENDER_ENUM = "L,P"
BOOL_ENUM = "TRUE,FALSE"

# Definisi tab: nama → header, seed, contoh, dropdown per kolom
SHEETS: dict[str, dict] = {
    "customers": {
        "headers": [
            "id", "phone_normalized", "name", "first_order_date", "created_at",
            "version", "branch", "order_count", "description", "age_range",
            "usia", "gender", "jenis_kelamin", "aliases",
            "branch_memberships", "is_followed_up", "followed_up_at",
        ],
        "dropdowns": {"branch": BRANCH_ENUM, "gender": GENDER_ENUM,
                      "is_followed_up": BOOL_ENUM},
    },
    "orders": {
        "headers": [
            "id", "customer_id", "order_date", "channel", "raw_phone_input",
            "created_at", "branch", "is_followed_up", "followed_up_at",
        ],
        "dropdowns": {"channel": CHANNEL_ENUM, "branch": BRANCH_ENUM,
                      "is_followed_up": BOOL_ENUM},
    },
    "settings": {
        "headers": ["key", "value"],
        "seed": [
            ["storeName", "Cabang Senopati"],
            ["activeDays", "30"],
            ["atRiskDays", "60"],
            ["waTemplate",
             "Halo {nama}, terima kasih sudah order di {toko}! 😊 Ada yang bisa kami bantu?"],
        ],
    },
    "branches": {
        "headers": ["id", "code", "name", "created_at", "address", "is_active"],
        "seed": [
            ["00000000-0000-0000-0000-000000000101", "CMH", "Cimahi (CMH)", "", "", "TRUE"],
            ["00000000-0000-0000-0000-000000000102", "BDG", "Bandung (BDG)", "", "", "TRUE"],
        ],
        "dropdowns": {"is_active": BOOL_ENUM},
    },
    "users": {
        "headers": [
            "id", "username", "display_name", "pin", "role", "created_at",
            "branch", "active", "status",
        ],
        "dropdowns": {"role": ROLE_ENUM, "branch": BRANCH_ENUM},
    },
    "customer_branches": {
        "headers": [
            "id", "customer_id", "branch", "first_order_date",
            "last_order_date", "order_count", "local_name", "notes",
            "created_at",
        ],
        "dropdowns": {"branch": BRANCH_ENUM},
    },
    "staging_cmh": {
        "headers": ["tanggal", "nama", "tipe_order", "no_telp", "status"],
        "dropdowns": {"tipe_order": CHANNEL_ENUM},
    },
    "staging_bdg": {
        "headers": ["tanggal", "nama", "tipe_order", "no_telp", "status"],
        "dropdowns": {"tipe_order": CHANNEL_ENUM},
    },
    # Tab import — header bahasa Indonesia sesuai src/utils/csvImport.ts
    "import_customers": {
        "headers": [
            "Nama", "No WhatsApp", "Cabang", "Gender", "Rentang Usia",
            "Tanggal Order Pertama", "Catatan",
        ],
        "sample": [["Budi Santoso", "081234567890", "CMH", "L", "26-35",
                    "2026-09-01", "Customer reguler"]],
    },
    "import_orders": {
        "headers": [
            "No WhatsApp", "Nama", "Tanggal Order", "Channel", "Cabang",
        ],
        "sample": [["081234567890", "Budi Santoso", "2026-09-01", "gofood", "CMH"]],
    },
}

PANDUAN = [
    ["CARA PAKAI", ""],
    ["1", "Upload file ini ke Google Drive, lalu buka dengan Google Sheets (Create Google Spreadsheet)."],
    ["2", "Share spreadsheet ke email service account sebagai Editor."],
    ["3", "Salin ID spreadsheet dari URL dan isi GOOGLE_SPREADSHEET_ID di .env.local."],
    ["4", "Tab _panduan ini BOLEH dihapus setelah konversi — tidak dibaca aplikasi."],
    ["", ""],
    ["ATURAN PENTING", ""],
    ["Jangan rename tab", "Nama tab harus persis: customers, orders, settings, branches, users, customer_branches, staging_cmh, staging_bdg."],
    ["Jangan ubah urutan kolom A-C customers", "sync-staging membaca per-index: A=id, B=phone_normalized, C=name."],
    ["Jangan ubah urutan kolom B-D orders", "sync-staging membaca per-index: B=customer_id, C=order_date, D=channel."],
    ["Nomor telepon = teks", "Selalu format kolom nomor HP sebagai Teks agar leading 0 tidak hilang."],
    ["aliases & branch_memberships", 'Diisi JSON string, contoh: [] atau ["CMH","BDG"]. Jangan diisi teks bebas.'],
    ["", ""],
    ["ENUM VALID", ""],
    ["channel (orders/staging)", CHANNEL_ENUM],
    ["branch", BRANCH_ENUM],
    ["role (users)", ROLE_ENUM],
    ["gender / jenis_kelamin", "L (Laki-laki), P (Perempuan)"],
    ["is_followed_up / is_active", "TRUE atau FALSE"],
    ["age_range", "0-17, 18-25, 26-35, 36-45, 46-55, 56-65, 66+"],
    ["", ""],
    ["FORMAT", ""],
    ["Tanggal", "YYYY-MM-DD (contoh: 2026-10-19)"],
    ["Timestamp", "ISO 8601 (contoh: 2026-10-19T10:30:00.000Z)"],
    ["", ""],
    ["TAB STAGING", ""],
    ["staging_cmh / staging_bdg", "Tempat paste data order harian per cabang (5 kolom). Setelah sync, kolom status otomatis diisi 'done'. Kolom tipe_order dipetakan ke channel."],
    ["", ""],
    ["TAB IMPORT", ""],
    ["import_customers / import_orders", "Acuan format upload massal lewat menu Import (Admin). Baris kuning = contoh, hapus sebelum dipakai."],
]


def build_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # hapus sheet default

    for name, spec in SHEETS.items():
        ws = wb.create_sheet(title=name)
        headers: list[str] = spec["headers"]
        text_cols = TEXT_COLUMNS.get(name, [])

        # Header
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        # Baris seed (data default aplikasi)
        for row_idx, seed_row in enumerate(spec.get("seed", []), start=2):
            for col_idx, value in enumerate(seed_row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = SEED_FILL
                if headers[col_idx - 1] in text_cols:
                    cell.number_format = "@"

        # Baris contoh (template import)
        for row_idx, sample_row in enumerate(
            spec.get("sample", []), start=2 + len(spec.get("seed", []))
        ):
            for col_idx, value in enumerate(sample_row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = SAMPLE_FILL
                if headers[col_idx - 1] in text_cols:
                    cell.number_format = "@"

        # Format teks untuk kolom telepon/id (input manual berikutnya juga)
        for col_name in text_cols:
            if col_name in headers:
                col_letter = get_column_letter(headers.index(col_name) + 1)
                for row_idx in range(2, 2001):
                    ws[f"{col_letter}{row_idx}"].number_format = "@"

        # Dropdown data validation
        for col_name, enum in spec.get("dropdowns", {}).items():
            if col_name not in headers:
                continue
            col_letter = get_column_letter(headers.index(col_name) + 1)
            dv = DataValidation(type="list", formula1=f'"{enum}"', allow_blank=True)
            dv.error = f"Nilai harus salah satu dari: {enum}"
            dv.errorTitle = "Nilai tidak valid"
            ws.add_data_validation(dv)
            dv.add(f"{col_letter}2:{col_letter}2000")

        # Freeze header + lebar kolom
        ws.freeze_panes = "A2"
        for col_idx, header in enumerate(headers, start=1):
            width = max(14, min(32, len(header) + 4))
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Tab _panduan (di luar allowlist aplikasi; boleh dihapus)
    ws = wb.create_sheet(title="_panduan")
    for row in PANDUAN:
        ws.append(row)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value and str(row[0].value).isupper():
            row[0].font = Font(bold=True)

    return wb


def verify(path: Path) -> None:
    """Baca ulang file hasil dan cocokkan nama tab + header terhadap SHEETS."""
    wb = load_workbook(path, read_only=True)
    expected_names = list(SHEETS.keys()) + ["_panduan"]
    assert wb.sheetnames == expected_names, (
        f"Sheet tidak sesuai.\n  Expected: {expected_names}\n  Actual:   {wb.sheetnames}"
    )
    for name, spec in SHEETS.items():
        ws = wb[name]
        actual = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert actual == spec["headers"], (
            f"Header {name} tidak sesuai.\n  Expected: {spec['headers']}\n  Actual:   {actual}"
        )
    wb.close()
    print("✅ Verifikasi lulus: semua nama tab & header sesuai skema aplikasi.")


def main() -> None:
    wb = build_workbook()
    wb.save(OUTPUT)
    print(f"✅ Template dibuat: {OUTPUT}")
    verify(OUTPUT)


if __name__ == "__main__":
    main()

